"""Import the supplied workbook without changing it; persist one seeded allocation.

Run with Python + openpyxl: python scripts/import-archives.py /path/to/清单.xlsx
Geometry follows 编号及图例.pptx on the two supplied 1448 × 1086 plans.
"""
import json
import random
import sys
from pathlib import Path

from openpyxl import load_workbook


def cabinet(identifier, source_id, x, y, width, height):
    return {"id": identifier, "sourceId": source_id, "bounds": [x, y, width, height]}


def make_halls():
    main = []
    west_x = [81, 137, 193, 248, 304, 360, 416, 474, 532, 590, 649]
    east_x = [758, 816, 874, 931, 989, 1046, 1103, 1159, 1215, 1270, 1324]
    for row, y in enumerate([337, 414, 492, 570, 650, 728], 1):
        for zone, positions in [("W", west_x), ("E", east_x)]:
            xs = positions[1:] if zone == "W" and row in (3, 4) else positions
            if zone == "E" and row in (3, 4):
                xs = positions[:-1]
            for index, x in enumerate(xs):
                col = len(xs) - index if zone == "W" else index + 1
                original_col = col
                # Slide 1 repeats E-02-08; normalize that row to eleven unique columns.
                if zone == "E" and row == 2 and col >= 9:
                    original_col = col - 1
                source_id = f"{zone}-{row:02}-{original_col:02}"
                main.append(cabinet(f"N-{zone}-{row:02}-{col:02}", source_id, x, y, 42, 65))
    for col, x in enumerate([649, 705, 758], 1):
        main.append(cabinet(f"N-M-01-{col:02}", f"M-01-{col:02}", x, 270, 42, 58))
    main.extend([
        cabinet("N-NW-01-01", "NW", 81, 289, 43, 44),
        cabinet("N-NE-01-01", "NE", 1323, 289, 43, 44),
    ])
    west = [cabinet("W-M-01-01", "M-01-01", 689, 348, 70, 86)]
    for row, y in enumerate([479, 644], 1):
        for col, x in enumerate([101, 177, 249, 324, 397, 472, 547, 621, 697,
                                 773, 848, 923, 998, 1072, 1147, 1220, 1295], 1):
            west.append(cabinet(f"W-{row:02}-{col:02}", f"L-{row:02}-{col:02}", x, y, 56, 83))
    return [
        {"id": "MainHall", "name": "主殿", "prefix": "N", "image": "maps/main-hall.png",
         "width": 1448, "height": 1086, "entrance": [724, 944],
         "walkable": [[64, 196, 1320, 688], [670, 880, 110, 72]],
         "obstacles": [[141, 194, 456, 55], [850, 194, 443, 55],
                       [83, 231, 35, 57], [705, 200, 39, 55], [1329, 231, 35, 57],
                       [64, 308, 8, 28], [64, 496, 8, 40], [64, 664, 8, 32],
                       [1374, 308, 10, 28], [1374, 496, 10, 40], [1374, 664, 10, 32]],
         "cabinets": main},
        {"id": "WestHall", "name": "西殿", "prefix": "W", "image": "maps/west-hall.png",
         "width": 1448, "height": 1086, "entrance": [724, 936],
         "walkable": [[66, 214, 1316, 638], [670, 848, 110, 96]],
         "obstacles": [[136, 212, 480, 70], [832, 212, 480, 70],
                       [66, 333, 7, 35], [66, 507, 8, 47], [66, 684, 7, 38],
                       [1374, 333, 8, 35], [1374, 507, 8, 47], [1374, 684, 8, 38]],
         "cabinets": west},
    ]


def main():
    source = Path(sys.argv[1])
    workbook = load_workbook(source, data_only=True)
    sheet = workbook["皇史宬馆藏档案"]
    halls = make_halls()
    slots = [(hall["id"], c["id"]) for hall in halls for c in hall["cabinets"]]
    random.Random(20261010).shuffle(slots)
    archives = []
    for row in range(3, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(1, 9)]
        if not values[0]:
            continue
        identifier, name, period, category, quantity, content, location, title = values
        link = sheet.cell(row, 8).hyperlink
        hall, cabinet_id = slots[len(archives)]
        archives.append({
            "id": str(identifier).zfill(3), "name": name, "period": period,
            "category": category, "quantity": quantity, "content": content,
            "currentLocation": location, "sourceTitle": title,
            "sourceUrl": link.target if link else None,
            "sourceRow": row, "hall": hall, "cabinetId": cabinet_id,
        })
    assert len({a["id"] for a in archives}) == len(archives)
    assert len({a["cabinetId"] for a in archives}) == len(archives)
    assert {a["hall"] for a in archives} == {"MainHall", "WestHall"}
    dataset = {"sourceWorkbook": source.name, "sourceSheet": sheet.title,
               "sourceDiagram": "编号及图例.pptx", "allocationSeed": 20261010,
               "allocationNote": "档案所在金匮为数字展示随机分配，不代表历史真实存放位置。",
               "halls": halls, "archives": archives}
    target = Path(__file__).resolve().parents[1] / "src/data/collection.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n")
    print(f"Imported {len(archives)} archives into {len(slots)} cabinets: {target}")


if __name__ == "__main__":
    main()
